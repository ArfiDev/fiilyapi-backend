"""IZN-B4d (earned_value 2/2) — rapor yüzeyi (QURR · günlük · panel · önizleme) maskesi.

🔴 Rapor yüzeyi TL TAŞIMAZ (ölçüldü): `CompositeCard.measure` yalnız `spent|earned|budget`
(adam-saat), `unit` payda iş tipinin ölçü birimi, `WarningOut.value` PF/saat/miktar. Kazanılmış /
harcanan / bütçe adam-saat, PF/SPI/ilerleme/oran, miktar, kişi sayısı, saat, tarih = `Hassas.yok`.
Bu yüzden HERHANGİ bir gizli kategori kümesi rapor yanıtını ve Excel'i DEĞİŞTİRMEZ; kişi adı
(`UserRef.full_name`) gizlenmez. Bekçi: modül `ZORUNLU_MODULLER`de.
"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.core.field_mask import Hassas, etiketler
from app.core.sayfalar import HiddenCategory
from app.modules.earned_value import schemas_reports
from tests._hassas_alan import rol_gizle

from .conftest import DAY
from .test_day_allocation import _body, _day
from .test_reports import _rep

pytestmark = pytest.mark.asyncio

_HEPSI = tuple(HiddenCategory)
#: Her çağrıda değişen üretim damgaları (değer değil zaman).
_DAMGALAR = {"generated_at"}


def _damgasiz(veri):  # noqa: ANN001, ANN202
    if isinstance(veri, dict):
        return {k: _damgasiz(v) for k, v in veri.items() if k not in _DAMGALAR}
    if isinstance(veri, list):
        return [_damgasiz(v) for v in veri]
    return veri


@pytest.fixture
async def dagitimli_gun(client, saha, santiye, boq, baseline, saha_gunu) -> None:
    body = _body(boq, saha_gunu["ali"], saha_gunu["veli"])
    assert (
        await client.put(_day(santiye, "/allocation"), headers=saha, json=body)
    ).status_code == 200


async def _uclar(client, headers, santiye) -> dict:  # noqa: ANN001
    """Rapor ailesinin JSON yanıtları + Excel hücreleri (tek rol, tek kümede)."""
    haftalik = await client.get(_rep(santiye, "/weekly"), headers=headers, params={"week": 1})
    gunluk = await client.get(
        _rep(santiye, "/daily"), headers=headers, params={"date": DAY.isoformat()}
    )
    panel = await client.get(
        f"/sites/{santiye.id}/earned-value/panel", headers=headers, params={"date": DAY.isoformat()}
    )
    onizleme = await client.get(
        f"/sites/{santiye.id}/earned-value/settings/preview",
        headers=headers,
        params={"date": DAY.isoformat()},
    )
    xlsx = await client.get(_rep(santiye, "/weekly.xlsx"), headers=headers, params={"week": 1})
    for yanit in (haftalik, gunluk, panel, onizleme, xlsx):
        assert yanit.status_code == 200, yanit.text
    hucreler = [
        [h.value for h in satir]
        for satir in load_workbook(BytesIO(xlsx.content)).active.iter_rows()
    ]
    return {
        "haftalik": _damgasiz(haftalik.json()),
        "gunluk": _damgasiz(gunluk.json()),
        "panel": _damgasiz(panel.json()),
        "onizleme": _damgasiz(onizleme.json()),
        "xlsx": hucreler,
    }


async def test_bayraksiz_rol_adam_saat_raporunu_gorur_pozitif_kontrol(
    client, sef, seeded_db, santiye, dagitimli_gun
):
    await rol_gizle(seeded_db, "site_chief")  # hiçbir kategori gizli değil

    uclar = await _uclar(client, sef, santiye)

    satir = next(r for r in uclar["haftalik"]["rows"] if r["h_earned_cum"] is not None)
    assert (Decimal(satir["h_earned_cum"]), Decimal(satir["i_spent_cum"])) == (10, 17)
    genel = next(k for k in uclar["gunluk"]["kpis"] if k["kind"] == "overall")
    assert Decimal(genel["earned_cum"]) == 10
    assert any(any(h is not None for h in satir_) for satir_ in uclar["xlsx"][2:])


async def test_gizli_roller_rapor_yanitini_ve_exceli_degistirmez(
    client, sef, seeded_db, santiye, dagitimli_gun
):
    await rol_gizle(seeded_db, "site_chief")
    acik = await _uclar(client, sef, santiye)

    for kume in (
        (HiddenCategory.maliyet_kar,),
        (HiddenCategory.sozlesme_fiyat, HiddenCategory.maliyet_kar),
        (HiddenCategory.tum_tutarlar,),
        _HEPSI,
    ):
        await rol_gizle(seeded_db, "site_chief", *kume)
        maskeli = await _uclar(client, sef, santiye)

        # Adam-saat TL değildir: null YOK, Excel hücreleri aynı, kişi adı görünür.
        assert maskeli == acik, kume
    assert acik["gunluk"]["kpis"][0]["budget_mhr"] is not None


async def test_rapor_semalarindaki_sayisal_alanlarin_hepsi_acikca_yok_etiketli():
    kategorili = []
    for ad in dir(schemas_reports):
        sema = getattr(schemas_reports, ad)
        if not (isinstance(sema, type) and hasattr(sema, "model_fields")):
            continue
        if sema.__module__ != schemas_reports.__name__:
            continue
        for alan, bilgi in sema.model_fields.items():
            kategoriler = {e for e in etiketler(bilgi) if e is not Hassas.yok}
            if kategoriler:
                kategorili.append(f"{ad}.{alan}")
    # TL taşıyan alan yok (ölçüldü): biri eklenirse bu test bilinçli güncellenmeli.
    assert kategorili == []
