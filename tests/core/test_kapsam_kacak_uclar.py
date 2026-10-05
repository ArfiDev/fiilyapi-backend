"""Maskenin KAÇAK UÇLARI — `BaseModel` DÖNMEYEN uçlar: export DAVRANIŞ bekçisi.

## Delik neydi

`GET /sites/{site_id}/boq/export` bir `Response` (xlsx baytları) döndürüyordu. Rota sarmalayıcısı
yalnız `BaseModel` sonuçlarını maskeler; dosya gövdesi maskesiz geçiyordu ve `boq` fiyatını gizleyen
rol, ekranda `—` gördüğü birim fiyatı ve tutarı Excel'den TAM DEĞERİYLE indiriyordu.

## Bu dosya NE ölçer

IZN-B4: YAPISAL bekçi (hangi uç maskeyi ELLE uyguluyor) `tests/core/test_hassas_alan_bekcisi.py`
içindeki `EXPORT_UCLARI` kaydına taşındı. Burada yalnız DAVRANIŞ kalır: gerçek rolle gerçek dosyayı
indirip hücrelere BAKAN `test_EXPORT_*` testleri. (Eski `_KapsamRotasi` tabanlı yapısal tarama
B4a'da konu kalmadığı için kalktı: hiçbir router eski köprüyü taşımıyor.)
"""

from decimal import Decimal
from io import BytesIO

import openpyxl

from tests.modules._boq import _auth, _group, _item, _login_with_access, _site

# --------------------------------------------------------------------------- #
# DAVRANIŞ BEKÇİSİ — gerçek rol, gerçek dosya
# --------------------------------------------------------------------------- #

_MIKTAR = Decimal("1240.000")
_BIRIM_FIYAT = Decimal("280.00")
#: `_MIKTAR * _BIRIM_FIYAT`. Elle yazılmıştır ki test, ölçtüğü formülü yeniden
#: hesaplayıp kendi kendini onaylamasın.
_TUTAR_METNI = "347200.00"

#: `app/modules/boq/export.py::COLUMN_HEADERS` sırası.
_SUTUN = {"poz": 1, "tarif": 2, "birim": 3, "miktar": 4, "birim_fiyat": 5, "tutar": 6}


async def _export_sayfasi(client, db_session, user_factory, project_factory, rol: str, eposta: str):
    """Gerçek rolle gerçek xlsx'i indirir ve sayfayı geri okur."""
    project = await project_factory(f"KACAK-{rol}")
    site = await _site(db_session, project, code=f"A-{rol[:6].upper()}")
    group = await _group(db_session, site)
    await _item(db_session, site, group, quantity=_MIKTAR, unit_price=_BIRIM_FIYAT)
    token = await _login_with_access(client, db_session, user_factory, rol, eposta)

    resp = await client.get(f"/sites/{site.id}/boq/export", headers=_auth(token))
    assert resp.status_code == 200, resp.text
    return site, token, openpyxl.load_workbook(BytesIO(resp.content)).active


def _hucreler(sayfa) -> list[str]:
    """Sayfadaki BOŞ OLMAYAN tüm hücrelerin metni — "hiçbir yerde yok" iddiası
    tek bir hücreye değil DOSYANIN TAMAMINA bakmalıdır."""
    return [str(h.value) for satir in sayfa.iter_rows() for h in satir if h.value is not None]


def _kalem_satiri(sayfa) -> int:
    """Başlık(1) + grup başlığı(2) + kalem(3). Sabit değil, aranarak bulunur."""
    for satir in range(1, sayfa.max_row + 1):
        if sayfa.cell(row=satir, column=_SUTUN["poz"]).value == "01.001":
            return satir
    raise AssertionError(f"Kalem satırı bulunamadı: {_hucreler(sayfa)}")


async def test_EXPORT_ADMIN_rolunde_TUM_degerler_dosyada(
    client, db_session, user_factory, project_factory
):
    """🔴 POZİTİF KONTROL — maske "herkesten gizle" hâline gelirse ya da export
    tamamen 403'e kapanırsa burası kırmızı olur. Aşağıdaki iki test tek başına
    `build_boq_workbook` boş dosya üretse de yeşil kalırdı."""
    _site_, _token, sayfa = await _export_sayfasi(
        client, db_session, user_factory, project_factory, "system_admin", "adm@kacak.co"
    )
    satir = _kalem_satiri(sayfa)

    assert sayfa.cell(row=satir, column=_SUTUN["miktar"]).value == "1240.000"
    assert sayfa.cell(row=satir, column=_SUTUN["birim_fiyat"]).value == "280.00"
    assert sayfa.cell(row=satir, column=_SUTUN["tutar"]).value == _TUTAR_METNI
    assert _TUTAR_METNI in _hucreler(sayfa), "GENEL TOPLAM satırı da dolu olmalı"


async def test_EXPORT_LIMITED_rolde_BIRIM_FIYAT_ve_TUTAR_dosyaya_YAZILMAZ(
    client, db_session, user_factory, project_factory
):
    """`site_chief` → `boq = view/limited`. Ekranda `—` gördüğü para, indirdiği
    dosyada da OLMAMALIDIR — aynı kapı (`boq:view`) iki farklı cevap veremez."""
    _site_, _token, sayfa = await _export_sayfasi(
        client, db_session, user_factory, project_factory, "site_chief", "sc@kacak.co"
    )
    satir = _kalem_satiri(sayfa)
    hucreler = _hucreler(sayfa)

    assert sayfa.cell(row=satir, column=_SUTUN["birim_fiyat"]).value is None, "BİRİM FİYAT SIZDI"
    assert sayfa.cell(row=satir, column=_SUTUN["tutar"]).value is None, "TÜREV TUTAR SIZDI"
    assert "280.00" not in hucreler, f"birim fiyat dosyanın BAŞKA bir hücresinde: {hucreler}"
    assert _TUTAR_METNI not in hucreler, f"GENEL TOPLAM sızdı: {hucreler}"
    # 🔴 "None" METNİ de aranır: `str(None)` yazan bir uygulama testin `is None`
    # iddiasını geçemez ama hücreye çöp basardı — ve kullanıcı onu veri sanardı.
    assert "None" not in hucreler, f"maskeli hücreye 'None' metni yazılmış: {hucreler}"
    # KİMLİK ve METRAJ DURUR: dosya şantiye şefi için kullanılabilir kalmalıdır;
    # 403 vermek onu işini yapamaz hâle getirirdi.
    assert sayfa.cell(row=satir, column=_SUTUN["poz"]).value == "01.001"
    assert sayfa.cell(row=satir, column=_SUTUN["miktar"]).value == "1240.000"


async def test_EXPORT_MUHASEBE_finance_KARSILIGI_YOK_hicbir_hucre_gizlenmez(
    client, db_session, user_factory, project_factory
):
    """IZN-B4: eski `finance` (metrajı gizler) yeni modelde KARŞILIKSIZDIR (IZN-PLAN §3: "finance
    karşılığı yok"). Muhasebe'nin `hidden_fields`ı B1 göçünde BOŞ türedi: dosyada metraj da
    birim fiyat da tutar da DOLUDUR. Gizleyen rol testi: `test_EXPORT_LIMITED_rolde_*`."""
    _site_, _token, sayfa = await _export_sayfasi(
        client, db_session, user_factory, project_factory, "accounting", "acc@kacak.co"
    )
    satir = _kalem_satiri(sayfa)

    assert sayfa.cell(row=satir, column=_SUTUN["miktar"]).value == "1240.000"
    assert sayfa.cell(row=satir, column=_SUTUN["birim_fiyat"]).value == "280.00"
    assert sayfa.cell(row=satir, column=_SUTUN["tutar"]).value == _TUTAR_METNI


async def test_EXPORT_dosyasi_EKRANLA_ayni_degerleri_tasir(
    client, db_session, user_factory, project_factory
):
    """🔴 ASIL INVARIANT: Excel ile ekran ASLA ayrışmaz.

    Sabit beklenen değerlere bakan testler, maskenin bir gün şemada değişmesi
    hâlinde (örn. `grand_total`ın `finance` davranışı) ikisinden yalnız birini
    yakalardı. Bu test iki yolu BİRBİRİNE karşı ölçer, sabite karşı değil.
    """
    site, token, sayfa = await _export_sayfasi(
        client, db_session, user_factory, project_factory, "site_chief", "sc2@kacak.co"
    )
    ekran = (await client.get(f"/sites/{site.id}/boq", headers=_auth(token))).json()
    kalem = ekran["groups"][0]["items"][0]
    satir = _kalem_satiri(sayfa)

    for alan, sutun in (
        ("quantity", "miktar"),
        ("unit_price", "birim_fiyat"),
        ("amount", "tutar"),
    ):
        assert sayfa.cell(row=satir, column=_SUTUN[sutun]).value == kalem[alan], (
            f"`{alan}` dosyada ekrandan FARKLI: "
            f"{sayfa.cell(row=satir, column=_SUTUN[sutun]).value!r} != {kalem[alan]!r}"
        )
