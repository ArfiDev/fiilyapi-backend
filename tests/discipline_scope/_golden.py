"""GOLDEN yardımcıları: UUID/zaman damgası normalizasyonu + xlsx hücre dökümü (DSC-B1).

SINIRLAR (bilinçli): `<ts>` ve ilk görünme sırasına göre `<uuid:N>` maskeleri TUTARLI biçimde
başka bir bilinmeyen kimliğe/zamana dönen yanıtı yakalamaz (ör. yanıttaki bilinmeyen bir
kimlik başka bir bilinmeyen kimlikle değişirse iki koşu aynı `<uuid:N>`i üretir). Yanıt
BAŞLIKLARI kıyaslanmaz (yalnız xlsx için content-type). Etiketlenmemiş kimlik sızıntısı
golden'da `<uuid:N>` olarak görünür; yeni uç eklerken dünya etiketlerini genişletin.
"""

from __future__ import annotations

import io
import json
import re
import uuid
from pathlib import Path

from openpyxl import load_workbook

GOLDEN_DIZINI = Path(__file__).parent / "golden"
_UUID = re.compile(r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}")
_ZAMAN = re.compile(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:\.\d+)?(?:Z|[+-]\d{2}:\d{2})?")


def normalize_metin(metin: str, etiketler: dict[uuid.UUID, str]) -> str:
    """Bilinen UUID → dünya etiketi; bilinmeyen → ilk görünme sırasıyla `<uuid:N>`;
    zaman damgası → `<ts>` (DB `now()` her koşuda farklıdır)."""
    sira: dict[str, str] = {}

    def degistir(m: re.Match[str]) -> str:
        ham = m.group(0)
        etiket = etiketler.get(uuid.UUID(ham))
        if etiket is not None:
            return etiket
        return sira.setdefault(ham, f"<uuid:{len(sira) + 1}>")

    return _UUID.sub(degistir, _ZAMAN.sub("<ts>", metin))


def normalize(govde: object, etiketler: dict[uuid.UUID, str]) -> object:
    ham = json.dumps(govde, ensure_ascii=False, sort_keys=True)
    return json.loads(normalize_metin(ham, etiketler))


def xlsx_dokumu(icerik: bytes, etiketler: dict[uuid.UUID, str]) -> list[list[object]]:
    """(sayfa, satır, sütun, değer) dökümü; boş hücreler atlanır."""
    kitap = load_workbook(io.BytesIO(icerik))
    dokum: list[list[object]] = []
    for sayfa in kitap.worksheets:
        for satir in sayfa.iter_rows():
            for hucre in satir:
                if hucre.value is None:
                    continue
                deger = normalize_metin(str(hucre.value), etiketler)
                dokum.append([sayfa.title, hucre.row, hucre.column, deger])
    return dokum


def golden_yolu(ad: str) -> Path:
    return GOLDEN_DIZINI / f"{ad}.json"


def yaz(ad: str, veri: object) -> None:
    GOLDEN_DIZINI.mkdir(exist_ok=True)
    golden_yolu(ad).write_text(
        json.dumps(veri, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )


def oku(ad: str) -> object:
    return json.loads(golden_yolu(ad).read_text(encoding="utf-8"))
