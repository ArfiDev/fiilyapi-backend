"""DSC-B3: kayıtlı bileşik kartlar (S3/S11), önizleme, Ü7, PUT /settings (Ü6) ve QURR xlsx."""

from __future__ import annotations

import io

from httpx import AsyncClient
from openpyxl import load_workbook

from app.modules.earned_value.report_router import qurr_workbook
from app.modules.earned_value.schemas_reports import QurrReport
from tests._disiplin_dunyasi import Dunya
from tests.discipline_scope._b3_dunya import KART_ADLARI, ayar_govdesi, kartlar
from tests.discipline_scope._b3_yardim import al, ev, haftalik

C1, C2, C3, C4 = KART_ADLARI
YETKI_YOK = "Bu işlem için yetkiniz yok"


def _kartlar(kartlar_: list[dict]) -> dict[str, dict]:
    return {k["name"]: k for k in kartlar_}


async def test_s3_kayitli_bilesik_kart_kurali_onizleme_ve_qurr(
    client: AsyncClient, dunya_b3: Dunya
) -> None:
    """C2 (I1+I2, karma) her tek-disiplinliye, C4 (d:none) HERKESE düşer; C1 yalnız Civil'e,
    C3 yalnız Elektrik'e; iki_disiplin C1+C2+C3 görür. Görünen kartın değeri atamasızla BİREBİR."""
    beklenen = {
        "atamasiz": {C1, C2, C3, C4},
        "civil": {C1},
        "elek": {C3},
        "iki_disiplin": {C1, C2, C3},
    }
    tam = _kartlar(
        (await al(client, dunya_b3, "/settings/preview", "atamasiz", {"date": "2026-05-06"}))[
            "composites"
        ]
    )
    qtam = _kartlar((await haftalik(client, dunya_b3, "atamasiz", 1))["composites"])
    for ad, adlar in beklenen.items():
        onizleme = _kartlar(
            (await al(client, dunya_b3, "/settings/preview", ad, {"date": "2026-05-06"}))[
                "composites"
            ]
        )
        qurr = _kartlar((await haftalik(client, dunya_b3, ad, 1))["composites"])
        assert set(onizleme) == adlar, ad
        assert set(qurr) == adlar, ad
        for kart in adlar:
            assert onizleme[kart] == tam[kart], (ad, kart)
            assert qurr[kart] == qtam[kart], (ad, kart)


async def test_s11_get_settings_ayni_kart_kurali(client: AsyncClient, dunya_b3: Dunya) -> None:
    beklenen = {
        "atamasiz": {C1, C2, C3, C4},
        "civil": {C1},
        "elek": {C3},
        "iki_disiplin": {C1, C2, C3},
        "pm_atamasiz": {C1, C2, C3, C4},
    }
    tam = await al(client, dunya_b3, "/settings", "atamasiz")
    for ad, adlar in beklenen.items():
        govde = await al(client, dunya_b3, "/settings", ad)
        assert {k["name"] for k in govde["composite_metrics"]} == adlar, ad
        # kartlar dışında ayar gövdesi aynı (takvim/bant/tolerans site düzeyi)
        assert {k: v for k, v in govde.items() if k != "composite_metrics"} == {
            k: v for k, v in tam.items() if k != "composite_metrics"
        }, ad
    for k in (await al(client, dunya_b3, "/settings", "civil"))["composite_metrics"]:
        assert k in tam["composite_metrics"]


async def test_u7_panel_yabanci_disiplin_olmayan_gibi(client: AsyncClient, dunya_b3: Dunya) -> None:
    d = dunya_b3
    params = {"date": "2026-05-06", "range": "4w"}
    yabanci = await client.get(
        ev(d, "/panel"),
        headers=d.baslik["elek"],
        params={**params, "discipline_id": f"d:{d.kab.id}"},
    )
    yok = await client.get(
        ev(d, "/panel"),
        headers=d.baslik["atamasiz"],
        params={**params, "discipline_id": f"d:{d.yabanci_kalem_kimligi}"},
    )
    assert yabanci.status_code == yok.status_code == 404
    assert yabanci.json() == yok.json()
    # d:none her kısıtlıya yabancı
    dnone = await client.get(
        ev(d, "/panel"), headers=d.baslik["civil"], params={**params, "discipline_id": "d:none"}
    )
    assert dnone.status_code == 404


async def test_u7_bilesik_onizleme_yabanci_kalem_olmayan_gibi(
    client: AsyncClient, dunya_b3: Dunya
) -> None:
    d = dunya_b3
    i2 = str(d.i["i2"].id)

    def params(pay: str) -> dict:
        return {
            "measure": "spent",
            "numerator_item_id": [pay],
            "denominator_item_id": i2,
            "date": "2026-05-06",
        }

    yabanci = await al(client, d, "/settings/preview/composite", "elek", params(str(d.i["i1"].id)))
    yok = await al(
        client, d, "/settings/preview/composite", "atamasiz", params(str(d.yabanci_kalem_kimligi))
    )
    assert yabanci == yok
    # yabancı PAYDA da görünmez (ad sızmaz)
    payda = await al(
        client,
        d,
        "/settings/preview/composite",
        "elek",
        {**params(i2), "denominator_item_id": str(d.i["i1"].id)},
    )
    assert payda["unit"] is None and payda["actual"] is None


async def test_put_settings_kisitliya_403_atamasiz_es_2xx(
    client: AsyncClient, dunya_b3: Dunya
) -> None:
    """F1: PM ve patron earned_value'da `_F` (settings YAZ kapısını geçer) → 403'ü DİSİPLİN
    kapısı verir; pozitif kontrol aynı rol: PM ↔ pm_atamasiz, patron ↔ yazar_atamasiz."""
    d = dunya_b3
    taban = await al(client, d, "/settings", "atamasiz")
    govde = ayar_govdesi(taban, kartlar(d), taban["tolerance_points"])
    url = ev(d, "/settings")
    for ad in ("civil", "elek", "iki_disiplin", "civil_yazar", "elek_yazar"):
        resp = await client.put(url, headers=d.baslik[ad], json=govde)
        assert resp.status_code == 403, (ad, resp.text)
        assert resp.json() == {"detail": YETKI_YOK}, ad
    for ad in ("pm_atamasiz", "yazar_atamasiz"):
        resp = await client.put(url, headers=d.baslik[ad], json=govde)
        assert resp.status_code == 200, (ad, resp.text)


def _hucreler(icerik: bytes) -> list[list[object]]:
    return [
        [h.value for h in satir]
        for s in load_workbook(io.BytesIO(icerik)).worksheets
        for satir in s.iter_rows()
    ]


async def test_xlsx_kisitli_hucre_dokumu_kisitli_json_ile_ayni(
    client: AsyncClient, dunya_b3: Dunya
) -> None:
    d = dunya_b3
    yabanci = {
        "civil": ("02.001", "Tuğla", "Elektrik", "Duvar"),
        "elek": ("01.001", "Beton", "Civil", "Betonarme"),
    }
    for ad in ("civil", "elek", "iki_disiplin"):
        json_ = await haftalik(client, d, ad, 1)
        resp = await client.get(
            ev(d, "/reports/weekly.xlsx"), headers=d.baslik[ad], params={"week": 1}
        )
        assert resp.status_code == 200
        beklenen = qurr_workbook(QurrReport.model_validate(json_))
        assert _hucreler(resp.content) == _hucreler(beklenen), ad
        metin = " ".join(
            str(v) for satir in _hucreler(resp.content) for v in satir if v is not None
        )
        for parca in yabanci.get(ad, ()):
            assert parca not in metin, (ad, parca)
        assert "Boya" not in metin and "03.001" not in metin, ad
