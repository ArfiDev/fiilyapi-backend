"""IZN-B5a madde 21a — mizan/bilanco/hesap plani/defter: YALNIZ bordronun besledigi hesaplar
(`335`, `361`, `730`) `maas_kisisel` ile de gizlenir; ara/genel toplamlar fail-closed."""

from __future__ import annotations

import ast
import io
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from app.core.sayfalar import HiddenCategory
from app.modules.accounting.bordro_hesaplari import (
    BORDRO_BESLENEN_HESAPLAR,
    bordro_hesabi_mi,
)
from app.modules.accounting.models import ChartAccountType
from tests._hassas_alan import rol_gizli
from tests.modules.accounting.conftest import _auth

H = HiddenCategory
_PAROLA = "parola1234"
_APP = Path(__file__).resolve().parents[3] / "app" / "modules"
_BORDRO_HESAPLARI = ("335", "361", "730")


# --- BEKCI: kume yedi *_POSTING_RULES sabitinden turetilir --------------------------------------


def _kurallar(dosya: str, sabit: str) -> set[str]:
    agac = ast.parse((_APP / dosya).read_text(encoding="utf-8"))
    for dugum in ast.walk(agac):
        if isinstance(dugum, ast.AnnAssign) and getattr(dugum.target, "id", None) == sabit:
            return {ast.literal_eval(e.elts[1]) for e in dugum.value.elts}  # type: ignore[union-attr]
    raise AssertionError(f"{dosya}: {sabit} yok")


def test_bekci_bordro_kumesi_otomatik_fis_kurallarindan_turetilir() -> None:
    bordro = _kurallar("payroll/posting.py", "PAYROLL_POSTING_RULES")
    digerleri = set().union(
        _kurallar("invoicing/posting.py", "INVOICE_POSTING_RULES"),
        _kurallar("progress_payments/posting.py", "PROGRESS_PAYMENT_POSTING_RULES"),
        _kurallar("subcontractor_progress_payments/posting.py", "SUBCONTRACTOR_POSTING_RULES"),
        _kurallar("treasury/posting.py", "PAYMENT_POSTING_RULES"),
        _kurallar("treasury/instruments/posting.py", "INSTRUMENT_POSTING_RULES"),
        _kurallar("equipment/rental_posting.py", "RENTAL_POSTING_RULES"),
    )
    assert frozenset(bordro - digerleri) == BORDRO_BESLENEN_HESAPLAR


def test_bordro_hesabi_alt_hesabi_da_kapsar() -> None:
    assert bordro_hesabi_mi("335") and bordro_hesabi_mi("335.01") and bordro_hesabi_mi("730")
    assert (
        not bordro_hesabi_mi("360") and not bordro_hesabi_mi("100") and not bordro_hesabi_mi("33")
    )


# --- Kurulum -------------------------------------------------------------------------------------


async def _kurulum(
    client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi, gizli, anahtar, bordro=True
):
    h = {}
    for kod, tur in (
        ("730", ChartAccountType.expense),
        ("335", ChartAccountType.liability),
        ("361", ChartAccountType.liability),
        ("100", ChartAccountType.asset),
        ("600", ChartAccountType.revenue),
    ):
        h[kod] = await hesap_fabrikasi(kod, name=f"Hesap {kod}", account_type=tur)
    # Elle (kaynaksiz) fis: hesap bazli kural kaynak tipine BAKMAZ.
    if bordro:
        await fis_fabrikasi(
            [(h["730"], "5000.00", "0"), (h["335"], "0", "4000.00"), (h["361"], "0", "1000.00")],
            entry_date=date(2026, 7, 5),
            description="Bordro tahakkuku",
        )
    await fis_fabrikasi(
        [(h["100"], "3000.00", "0"), (h["600"], "0", "3000.00")],
        entry_date=date(2026, 7, 6),
        description="Satis",
    )
    await seeded_db.flush()
    rol = await rol_gizli(seeded_db, f"mizan_{anahtar}", gizli)
    e_posta = f"{anahtar}@b5a.co"
    await user_factory(email=e_posta, password=_PAROLA, role_key=rol.key)
    yanit = await client.post("/auth/login", json={"email": e_posta, "password": _PAROLA})
    return _auth(yanit.json()["access_token"])


_ALANLAR = (
    "opening_debit", "opening_credit", "period_debit", "period_credit", "closing_debit",
    "closing_credit",
)  # fmt: skip


async def _mizan(client, basliklar):
    yanit = await client.get("/trial-balance?year=2026&month=7", headers=basliklar)
    assert yanit.status_code == 200
    return yanit.json()


async def test_mizan_maas_kisisel_gizliyken_bordro_satirlari_ve_genel_toplam_NULL(
    client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi
):
    basliklar = await _kurulum(
        client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi, {H.maas_kisisel}, "g1"
    )
    govde = await _mizan(client, basliklar)
    satirlar = {s["account_code"]: s for s in govde["rows"]}

    for kod in _BORDRO_HESAPLARI:
        assert all(satirlar[kod][a] is None for a in _ALANLAR), kod
    # POZITIF KONTROL: bordro disi hesap etkilenmez.
    assert satirlar["100"]["period_debit"] == "3000.00"
    assert satirlar["600"]["period_credit"] == "3000.00"
    # Genel toplam bordro satirlarini icerir → fail-closed (cikarma yoluyla turetme kapali).
    assert all(govde["totals"][a] is None for a in _ALANLAR)


async def test_mizan_kategorisi_acik_rol_tam_gorur(
    client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi
):
    basliklar = await _kurulum(
        client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi, set(), "a1"
    )
    govde = await _mizan(client, basliklar)
    satirlar = {s["account_code"]: s for s in govde["rows"]}
    assert satirlar["335"]["period_credit"] == "4000.00"
    assert satirlar["730"]["period_debit"] == "5000.00"
    assert govde["totals"]["period_debit"] == "8000.00"


async def test_mizan_bordro_hareketi_yoksa_genel_toplam_gizli_rolde_de_gorunur(
    client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi
):
    """Ince sinir: bordro hesaplarinda hareket yoksa toplam bordro icermez → gereksiz maske YOK."""
    basliklar = await _kurulum(
        client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi, {H.maas_kisisel}, "g2",
        bordro=False,
    )  # fmt: skip
    govde = await _mizan(client, basliklar)
    assert govde["totals"]["period_debit"] == "3000.00"
    assert {s["account_code"] for s in govde["rows"]} == {"100", "600"}


async def test_mizan_sema_alani_eklenmedi_openapi_degismez() -> None:
    from app.modules.accounting.reports_schemas import TrialBalanceTotals
    from app.modules.accounting.schemas import LedgerResponse

    assert "_bordro_dahil" not in TrialBalanceTotals.model_fields
    assert "_bordro_hesabi" not in LedgerResponse.model_fields


async def test_mizan_xlsx_bordro_tutarlarini_sizdirmaz(
    client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi
):
    gizli = await _kurulum(
        client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi, {H.maas_kisisel}, "g3"
    )
    yanit = await client.get("/trial-balance/export.xlsx?year=2026&month=7", headers=gizli)
    assert yanit.status_code == 200
    hucreler = [
        str(h.value)
        for satir in load_workbook(io.BytesIO(yanit.content)).active.iter_rows()
        for h in satir
        if h.value is not None
    ]
    metin = " ".join(hucreler)
    for tutar in ("4000", "5000", "1000", "8000"):
        assert tutar not in metin, tutar
    assert "3000" in metin  # bordro disi satir korunur (pozitif kontrol)


async def test_bilanco_bordro_kalemi_ara_toplam_ve_taraf_toplami_gizlenir(
    client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi
):
    gizli = await _kurulum(
        client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi, {H.maas_kisisel}, "g4"
    )
    govde = (await client.get("/balance-sheet?as_of=2026-07-31", headers=gizli)).json()

    def kalemler(taraf):
        return [(b, k) for b in taraf["sections"] for k in b["lines"]]

    bordro_kalemleri = [
        (b, k)
        for taraf in (govde["assets"], govde["liabilities"])
        for b, k in kalemler(taraf)
        if any(c.split(".")[0] in _BORDRO_HESAPLARI for c in k["account_codes"])
    ]
    assert bordro_kalemleri, "bordro hesabi bir kaleme dusmeli"
    for bolum, kalem in bordro_kalemleri:
        assert kalem["amount"] is None, kalem["key"]
        assert bolum["subtotal"] is None, bolum["key"]
    assert govde["liabilities"]["total"] is None
    # POZITIF KONTROL: bordro icermeyen kalem (nakit) ve aktif taraf toplami gorunur.
    nakit = [k for _, k in kalemler(govde["assets"]) if "100" in k["account_codes"]]
    assert nakit and nakit[0]["amount"] == "3000.00"
    assert govde["assets"]["total"] == "3000.00"


async def test_bilanco_kategorisi_acik_rol_tam_gorur(
    client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi
):
    acik = await _kurulum(
        client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi, set(), "a2"
    )
    govde = (await client.get("/balance-sheet?as_of=2026-07-31", headers=acik)).json()
    assert govde["liabilities"]["total"] is not None
    assert all(
        k["amount"] is not None for b in govde["liabilities"]["sections"] for k in b["lines"]
    )


async def test_hesap_plani_bakiyesi_bordro_hesabinda_gizlenir(
    client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi
):
    gizli = await _kurulum(
        client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi, {H.maas_kisisel}, "g5"
    )
    govde = (await client.get("/chart-of-accounts?limit=200", headers=gizli)).json()
    ogeler = {o["code"]: o for o in govde["items"]}
    for kod in _BORDRO_HESAPLARI:
        assert ogeler[kod]["balance"] is None, kod
    assert ogeler["100"]["balance"] == "3000.00"


async def test_defter_bordro_hesabi_satirinda_tutar_ve_koşan_bakiye_gizlenir(
    client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi
):
    gizli = await _kurulum(
        client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi, {H.maas_kisisel}, "g6"
    )
    govde = (await client.get("/journal?year=2026&month=7&limit=200", headers=gizli)).json()
    satirlar = govde["items"]
    for s in satirlar:
        if s["account_code"] in _BORDRO_HESAPLARI:
            assert s["debit"] is None and s["credit"] is None and s["running_balance"] is None
        else:
            assert s["debit"] is not None or s["credit"] is not None


async def test_defter_bordro_hesabi_filtresinde_devreden_bakiye_gizlenir(
    client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi
):
    gizli = await _kurulum(
        client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi, {H.maas_kisisel}, "g7"
    )
    from sqlalchemy import select

    from app.modules.accounting.models import ChartAccount

    async def kimlik(kod):
        return (
            await seeded_db.execute(select(ChartAccount.id).where(ChartAccount.code == kod))
        ).scalar_one()

    # Agustos penceresi: devreden = Temmuz birikimi.
    bordro = (
        await client.get(
            f"/journal?year=2026&month=8&account_id={await kimlik('335')}", headers=gizli
        )
    ).json()
    nakit = (
        await client.get(
            f"/journal?year=2026&month=8&account_id={await kimlik('100')}", headers=gizli
        )
    ).json()
    assert bordro["carried_balance"] is None
    assert nakit["carried_balance"] == "3000.00"  # pozitif kontrol


# --- Gelir tablosu + nakit akis (cürütücü kanit senaryosu) ---------------------------------------


async def _kanit_kurulum(
    client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi, gizli, anahtar, bordro=True
):
    h = {}
    for kod, tur in (
        ("730", ChartAccountType.expense),
        ("335", ChartAccountType.liability),
        ("100", ChartAccountType.asset),
        ("600", ChartAccountType.revenue),
    ):
        h[kod] = await hesap_fabrikasi(kod, name=f"Hesap {kod}", account_type=tur)
    await fis_fabrikasi(
        [(h["100"], "9000.00", "0"), (h["600"], "0", "9000.00")], entry_date=date(2026, 7, 2)
    )
    if bordro:
        await fis_fabrikasi(
            [(h["730"], "5000.00", "0"), (h["335"], "0", "5000.00")], entry_date=date(2026, 7, 5)
        )
        await fis_fabrikasi(
            [(h["335"], "4000.00", "0"), (h["100"], "0", "4000.00")], entry_date=date(2026, 7, 6)
        )
    await seeded_db.flush()
    rol = await rol_gizli(seeded_db, f"kanit_{anahtar}", gizli)
    await user_factory(email=f"{anahtar}@b5a.co", password=_PAROLA, role_key=rol.key)
    yanit = await client.post(
        "/auth/login", json={"email": f"{anahtar}@b5a.co", "password": _PAROLA}
    )
    return _auth(yanit.json()["access_token"])


def _kalemler(govde):
    return [(b, k) for b in govde["sections"] for k in b["lines"]]


async def test_gelir_tablosu_bordro_satiri_ara_toplam_ve_toplamlar_gizlenir(
    client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi
):
    gizli = await _kanit_kurulum(
        client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi, {H.maas_kisisel}, "i1"
    )
    govde = (await client.get("/income-statement?year=2026&month=7", headers=gizli)).json()
    bordro = [(b, k) for b, k in _kalemler(govde) if "730" in k["account_codes"]]
    assert bordro
    for bolum, kalem in bordro:
        assert kalem["amount"] is None and bolum["subtotal"] is None
    assert govde["total_expense"] is None and govde["period_profit"] is None
    assert govde["total_revenue"] == "9000.00"  # pozitif kontrol: gelir bagimsiz


async def test_gelir_tablosu_kategorisi_acik_rol_tam_gorur_ve_bordrosuz_donemde_maske_yok(
    client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi
):
    acik = await _kanit_kurulum(
        client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi, set(), "i2"
    )
    govde = (await client.get("/income-statement?year=2026&month=7", headers=acik)).json()
    assert govde["total_expense"] == "5000.00" and govde["period_profit"] == "4000.00"


async def test_gelir_tablosu_bordro_hareketi_yoksa_gizli_rolde_de_gorunur(
    client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi
):
    gizli = await _kanit_kurulum(
        client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi, {H.maas_kisisel}, "i3",
        bordro=False,
    )  # fmt: skip
    govde = (await client.get("/income-statement?year=2026&month=7", headers=gizli)).json()
    assert govde["period_profit"] == "9000.00" and Decimal(govde["total_expense"]) == 0


async def test_nakit_akisi_bordro_satiri_bolum_ve_toplamlar_gizlenir(
    client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi
):
    gizli = await _kanit_kurulum(
        client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi, {H.maas_kisisel}, "n1"
    )
    govde = (await client.get("/cash-flow-statement?year=2026&month=7", headers=gizli)).json()
    bordro = [(b, k) for b, k in _kalemler(govde) if "335" in k["account_codes"]]
    assert bordro, "335 bir kaleme dusmeli"
    for bolum, kalem in bordro:
        assert kalem["amount"] is None and bolum["subtotal"] is None
    assert govde["net_change"] is None and govde["closing_cash"] is None
    assert all(n["closing_cash"] is None for n in govde["monthly_cash"])
    assert govde["opening_cash"] is not None  # onceki donem birikimi, bu donem akisi degil


async def test_nakit_akisi_kategorisi_acik_rol_tam_gorur_bordrosuz_donemde_maske_yok(
    client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi
):
    acik = await _kanit_kurulum(
        client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi, set(), "n2"
    )
    govde = (await client.get("/cash-flow-statement?year=2026&month=7", headers=acik)).json()
    assert govde["net_change"] is not None and govde["closing_cash"] == "5000.00"


async def test_nakit_akisi_bordro_hareketi_yoksa_gizli_rolde_de_gorunur(
    client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi
):
    gizli = await _kanit_kurulum(
        client, seeded_db, user_factory, hesap_fabrikasi, fis_fabrikasi, {H.maas_kisisel}, "n3",
        bordro=False,
    )  # fmt: skip
    govde = (await client.get("/cash-flow-statement?year=2026&month=7", headers=gizli)).json()
    assert govde["closing_cash"] == "9000.00"
