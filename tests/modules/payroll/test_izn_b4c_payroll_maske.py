"""IZN-B4c — bordro modülü hassas alan maskesi (`maas_kisisel`).

Gizleyen rol: `maas_kisisel` gizli özel rol (`gizli_headers`; `hr_manager` madde 20 ile artık
GÖRÜR). Gizlemeyen rol: `hr_manager` (`ik_headers`, bordronun gerçek kullanıcısı). Oranlar, vergi
dilimi sınırları ve gün `Hassas.yok`tur: gizli
rolde de GÖRÜNÜR (şirket geneli mevzuat parametresi, kişisel veri değil).
"""

from io import BytesIO

import openpyxl
import pytest

from app.modules.payroll.export import HEADER_ROW, SHEET_TITLE

pytestmark = pytest.mark.asyncio

SATIR_TUTARLARI = (
    "gross_amount",
    "deduction_amount",
    "net_amount",
    "bank_amount",
    "cash_amount",
    "tax_base_amount",
    "cumulative_tax_base",
    "income_tax_amount",
)
OZET_TUTARLARI = (
    "net_total",
    "bank_total",
    "cash_total",
    "gross_total",
    "sgk_employer_total",
    "total_employer_cost",
)


@pytest.fixture
async def hesaplanmis(client, ik_headers, donem, dort_tip):
    resp = await client.post(f"/payroll/periods/{donem.id}/compute", headers=ik_headers)
    assert resp.status_code == 200, resp.text
    return donem


def _satirlar(detay: dict) -> dict[str, dict]:
    return {s["personnel_name"]: s for b in detay["sections"] for s in b["lines"]}


async def _detay(client, headers, donem) -> dict:
    resp = await client.get(f"/payroll/periods/{donem.id}", headers=headers)
    assert resp.status_code == 200, resp.text
    return resp.json()


async def test_gizlemeyen_rol_bordro_tutarlarini_gorur(client, ik_headers, hesaplanmis):
    """Pozitif kontrol: maskeleme bağlanmamış olsaydı bu da yeşil kalırdı — gizli rol testi
    ancak bu sayede anlam taşır."""
    detay = await _detay(client, ik_headers, hesaplanmis)
    ayse = _satirlar(detay)["Ayşe Demir"]
    assert ayse["gross_amount"] == "9000.00"
    assert ayse["net_amount"] == "7650.00"
    assert ayse["days"] is not None
    assert detay["summary"]["net_total"] is not None
    assert detay["summary"]["gross_total"] is not None


async def test_gizli_rolde_satir_ve_ozet_tutarlari_null(client, gizli_headers, hesaplanmis):
    detay = await _detay(client, gizli_headers, hesaplanmis)
    for satir in _satirlar(detay).values():
        for alan in SATIR_TUTARLARI:
            assert satir[alan] is None, alan
        # Gün `yok`: gizlenmez.
    assert _satirlar(detay)["Ayşe Demir"]["days"] == "5.0"
    for alan in OZET_TUTARLARI:
        assert detay["summary"][alan] is None, alan
    # Sayaç ve yüzde `yok`/etiketsiz sayaçtır: maskelenmez.
    assert detay["summary"]["line_count"] == 5
    assert "bank_pct" in detay["summary"]


async def test_gizli_rolde_donem_listesi_tutarlari_null(client, gizli_headers, hesaplanmis):
    resp = await client.get("/payroll/periods", headers=gizli_headers)
    assert resp.status_code == 200, resp.text
    satir = resp.json()["items"][0]
    for alan in ("gross_total", "sgk_employer_total", "net_total", "total_cost"):
        assert satir[alan] is None, alan
    assert satir["personnel_count"] == 5


async def test_gizli_rolde_sgk_ozeti_tutarlari_null_sayaclar_dolu(
    client, ik_headers, gizli_headers, hesaplanmis
):
    yol = f"/payroll/periods/{hesaplanmis.id}/sgk-summary"
    acik = (await client.get(yol, headers=ik_headers)).json()
    gizli = (await client.get(yol, headers=gizli_headers)).json()
    assert acik["sgk_premium_total"] is not None
    for alan in (
        "sgk_base_total",
        "sgk_premium_total",
        "unemployment_total",
        "employee_deduction_total",
        "employer_burden_total",
        "sgk_payable_total",
    ):
        assert gizli[alan] is None, alan
    assert gizli["declared_personnel_count"] == acik["declared_personnel_count"]


async def test_oranlar_ve_vergi_dilimleri_gizli_rolde_gorunur(
    client, gizli_headers, hesaplanmis, oranlar
):
    oran = (await client.get("/payroll/rates", headers=gizli_headers)).json()["items"]
    assert oran and all(o["sgk_employee_pct"] is not None for o in oran)
    dilim = (await client.get("/payroll/tax-brackets", headers=gizli_headers)).json()["items"]
    assert dilim and any(d["upper_bound"] is not None for d in dilim)


async def test_gizli_rolde_dolu_tutarla_PATCH_403(client, gizli_headers, hesaplanmis):
    detay = await _detay(client, gizli_headers, hesaplanmis)
    satir_id = _satirlar(detay)["Ayşe Demir"]["id"]
    resp = await client.patch(
        f"/payroll/lines/{satir_id}", json={"gross_amount": "10000.00"}, headers=gizli_headers
    )
    assert resp.status_code == 403, resp.text


async def test_gizlemeyen_rolde_PATCH_calisir(client, ik_headers, hesaplanmis):
    detay = await _detay(client, ik_headers, hesaplanmis)
    satir_id = _satirlar(detay)["Ayşe Demir"]["id"]
    resp = await client.patch(
        f"/payroll/lines/{satir_id}", json={"gross_amount": "10000.00"}, headers=ik_headers
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["gross_amount"] == "10000.00"


def _hucreler(icerik: bytes, ilk_satir: int = 1) -> list:
    sheet = openpyxl.load_workbook(BytesIO(icerik)).worksheets[0]
    return [
        sheet.cell(row=r, column=c).value
        for r in range(ilk_satir, sheet.max_row + 1)
        for c in range(1, sheet.max_column + 1)
    ]


async def test_gizli_rolde_donem_excel_tutar_hucreleri_bos(
    client, ik_headers, gizli_headers, hesaplanmis
):
    yol = f"/payroll/periods/{hesaplanmis.id}/export"
    acik = await client.get(yol, headers=ik_headers)
    gizli = await client.get(yol, headers=gizli_headers)
    assert gizli.status_code == 200, gizli.text
    assert "9000.00" in _hucreler(acik.content)
    hucre = _hucreler(gizli.content)
    assert "9000.00" not in hucre and "7650.00" not in hucre
    assert "None" not in hucre
    assert "Ayşe Demir" in hucre  # kimlik/ad gizlenmez
    sheet = openpyxl.load_workbook(BytesIO(gizli.content))[SHEET_TITLE]
    assert sheet.cell(row=4, column=2).value is None  # "Toplam Net Ödenecek" değeri boş
    assert sheet.cell(row=HEADER_ROW, column=4).value == "Brüt"  # başlıklar kalır


async def test_gizli_rolde_gecmis_excel_tutar_hucreleri_bos(
    client, ik_headers, gizli_headers, hesaplanmis
):
    acik = await client.get("/payroll/periods/export.xlsx", headers=ik_headers)
    gizli = await client.get("/payroll/periods/export.xlsx", headers=gizli_headers)
    assert gizli.status_code == 200, gizli.text
    assert any(h is not None and h != "—" and "." in str(h) for h in _hucreler(acik.content, 2))
    satir = openpyxl.load_workbook(BytesIO(gizli.content)).worksheets[0]
    assert [satir.cell(row=2, column=c).value for c in (3, 4, 5, 6)] == [None] * 4
    assert satir.cell(row=2, column=2).value == "5"  # çalışan sayısı sayaçtır
