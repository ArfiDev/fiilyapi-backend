"""IZN-B5a madde 19/21b — denetim gunlugu metinlerindeki TUTAR sizintisi (Politika A).

Kural: tutar tasiyan olayin metni, okuyucunun ilgili kategorisi gizliyse sabit ifadeyle degisir;
tur/aktor/zaman kalir; kategorisi acik rol TAM metni gorur; Excel ayni kurala uyar.
"""

from __future__ import annotations

import ast
import io
import re
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import delete

from app.core.field_mask import Hassas
from app.core.sayfalar import HiddenCategory
from app.modules.audit import messages
from app.modules.audit.gizli_tutar import (
    GIZLI_TUTAR_METNI,
    TUTAR_DEGIL,
    TUTAR_OLAYLARI,
    detay_maskele,
    gizli_olaylar,
)
from app.modules.audit.models import AuditAction, AuditLog
from tests._hassas_alan import rol_gizli

H = HiddenCategory
_PAROLA = "parola1234"
_MESAJ_DIZINI = Path(messages.__file__).parent

# --- Ornek metinler: HER tutarli olay icin gercek mesaj fonksiyonu -------------------------------
_TUTAR = Decimal("1250.00")
ORNEKLER: dict[str, str] = {
    "work_item_price_updated": messages.work_item_price_updated(
        "15.100.1001", "Beton", "m3", Decimal("100"), Decimal("1250.00")
    ),
    "progress_payment_deleted": messages.progress_payment_deleted("Kule", 3, "Taslak", _TUTAR),
    "subcontractor_progress_payment_deleted": messages.subcontractor_progress_payment_deleted(
        "Kule", "Usta A.Ş.", 2, "Taslak", _TUTAR
    ),
    "payroll_period_paid": messages.payroll_period_paid(2026, 7, 12, _TUTAR),
    "sale_installment_paid": messages.sale_installment_paid("Kule", "A-12", "Taksit 1", _TUTAR),
}
TUTARSIZ = messages.work_item_updated("15.100.1001", "Beton", "m3")

# kategori bayragi -> o bayrakla gizlenmesi GEREKEN olaylar
BEKLENEN_GIZLI: dict[H, set[str]] = {
    H.sozlesme_fiyat: {
        "work_item_price_updated",
        "progress_payment_deleted",
        "subcontractor_progress_payment_deleted",
    },
    H.maliyet_kar: {"subcontractor_progress_payment_deleted"},
    H.maas_kisisel: {"payroll_period_paid"},
    H.satis_alici: {"sale_installment_paid"},
    H.banka_kasa: set(),
    H.tum_tutarlar: set(ORNEKLER),
}


# --- BEKCI: tablo mesaj paketiyle birebir -------------------------------------------------------

_PARA_ADI_ANAHTAR = frozenset({"amount", "tutar", "fiyat", "bakiye", "balance", "gross"})
#: `int` SAYAC olabilen adlar (`total_steps`, `price_updated`): yalniz int DISINDA para sayilir.
_PARA_ADI_INT_DISI = frozenset({"price", "total"})


def _para_parametresi_mi(parametre: ast.arg) -> bool:
    belirtec = set(parametre.arg.split("_"))
    if belirtec & _PARA_ADI_ANAHTAR:
        return True
    tip = ast.unparse(parametre.annotation) if parametre.annotation else ""
    return bool(belirtec & _PARA_ADI_INT_DISI) and tip != "int"


_APP_DIZINI = _MESAJ_DIZINI.parents[2]  # app/


def _mesaj_dosyalari() -> list[Path]:
    """`audit/messages/*.py` + repodaki HER `*audit_messages*.py` (ornek: earned_value)."""
    dosyalar = [f for f in _MESAJ_DIZINI.glob("*.py") if f.name != "__init__.py"]
    dosyalar += [f for f in _APP_DIZINI.rglob("*audit_messages*.py") if f not in dosyalar]
    return dosyalar


def _para_ifadesi_mi(kaynak: str) -> bool:
    return "₺" in kaynak or ",.2f" in kaynak or bool(re.search(r"\bTL\b", kaynak))


def _para_gibi_fonksiyonlar(kaynak: str) -> set[str]:
    """Kaynaktaki para bicimlendiren/tasiyan fonksiyonlar (AST): Decimal/object parametre, para adli
    parametre ya da govdede TL/₺/`,.2f` bicimi."""
    bulunan: set[str] = set()
    for dugum in ast.parse(kaynak).body:
        if not isinstance(dugum, ast.FunctionDef):
            continue
        govde = ast.get_source_segment(kaynak, dugum) or ""
        parametreler = dugum.args.args + dugum.args.kwonlyargs
        tipler = " ".join(ast.unparse(p.annotation) for p in parametreler if p.annotation)
        if (
            "Decimal" in tipler
            or "object" in tipler
            or any(_para_parametresi_mi(p) for p in parametreler)
            or "TL" in govde.split('"""')[-1]
            or "₺" in govde
            or ",.2f" in govde
        ):
            bulunan.add(dugum.name)
    return bulunan


def _para_gibi_satir_ici_detaylar(kaynak: str) -> list[int]:
    """`record_audit(..., detail=<f-string/format/%/+>)` cagrilarinda PARA bicimi tasiyanlarin
    satir numarasi. Mesaj fonksiyonuna giden cagrilar (`detail=messages.x(...)`) bu kapsamda degil:
    onlari mesaj bekcisi izler."""
    bulunan: list[int] = []
    for dugum in ast.walk(ast.parse(kaynak)):
        if not isinstance(dugum, ast.Call):
            continue
        ad = getattr(dugum.func, "attr", getattr(dugum.func, "id", ""))
        if ad != "record_audit":
            continue
        for kw in dugum.keywords:
            if kw.arg == "detail" and isinstance(kw.value, ast.JoinedStr | ast.BinOp):
                parca = ast.get_source_segment(kaynak, kw.value) or ""
                if _para_ifadesi_mi(parca) or any(
                    t in parca.lower() for t in ("amount", "tutar", "price", "total", "fiyat")
                ):
                    bulunan.append(dugum.lineno)
    return bulunan


def _para_gibi_mesajlar() -> set[str]:
    bulunan: set[str] = set()
    for dosya in _mesaj_dosyalari():
        bulunan |= _para_gibi_fonksiyonlar(dosya.read_text(encoding="utf-8"))
    return bulunan


def test_bekci_kapsami_audit_messages_dosyalarini_ve_messages_dizinini_icerir() -> None:
    adlar = {d.name for d in _mesaj_dosyalari()}
    assert "audit_messages.py" in adlar and "payroll.py" in adlar


def test_bekci_satir_ici_audit_detaylarinda_para_yok() -> None:
    kirli = {
        str(d.relative_to(_APP_DIZINI)): satirlar
        for d in _APP_DIZINI.rglob("*.py")
        if (satirlar := _para_gibi_satir_ici_detaylar(d.read_text(encoding="utf-8")))
    }
    assert not kirli, f"satir ici para iceren record_audit detail: {kirli}"


def test_bekci_sentetik_ihlalleri_yakalar() -> None:
    """POZITIF KONTROL: bekci kendi kapsaminda sentetik kaynakta FIILEN ateslenir."""
    mesaj = "def x_paid(label: str, amount: int) -> str:\n    return f'{label} {amount:,.2f} TL'\n"
    assert _para_gibi_fonksiyonlar(mesaj) == {"x_paid"}
    temiz = "def x_named(label: str) -> str:\n    return f'{label}'\n"
    assert _para_gibi_fonksiyonlar(temiz) == set()
    satir_ici = "record_audit(s, detail=f'Odeme {tutar:,.2f} TL', action=a)\n"
    assert _para_gibi_satir_ici_detaylar(satir_ici) == [1]
    assert _para_gibi_satir_ici_detaylar("record_audit(s, detail=f'Ad {ad}', action=a)\n") == []
    assert _para_gibi_satir_ici_detaylar("record_audit(s, detail=messages.x(a), action=a)\n") == []


def test_bekci_tutar_bicimlendiren_her_mesaj_tabloda_ya_da_gerekceli_istisnada() -> None:
    tabloda = {o.mesaj for o in TUTAR_OLAYLARI}
    kayitsiz = _para_gibi_mesajlar() - tabloda - set(TUTAR_DEGIL)
    assert not kayitsiz, (
        f"TUTAR_OLAYLARI/TUTAR_DEGIL'de yok (yeni tutarli mesaj?): {sorted(kayitsiz)}"
    )


def test_bekci_tablo_bayat_kayit_ve_celiski_icermez() -> None:
    mevcut = {
        d.name
        for f in _mesaj_dosyalari()
        for d in ast.parse(f.read_text(encoding="utf-8")).body
        if isinstance(d, ast.FunctionDef)
    }
    tabloda = {o.mesaj for o in TUTAR_OLAYLARI}
    assert tabloda <= mevcut, f"bayat tablo kaydi: {sorted(tabloda - mevcut)}"
    assert set(TUTAR_DEGIL) <= mevcut, f"bayat istisna: {sorted(set(TUTAR_DEGIL) - mevcut)}"
    assert not tabloda & set(TUTAR_DEGIL)
    assert all(o.kategoriler and Hassas.yok not in o.kategoriler for o in TUTAR_OLAYLARI)


def test_bekci_onekler_gercek_mesajlarla_eslesir_ve_tutarsiz_kardes_eslesmez() -> None:
    assert set(ORNEKLER) == {o.mesaj for o in TUTAR_OLAYLARI}, "her olay icin ornek yazilmali"
    for olay in TUTAR_OLAYLARI:
        assert olay.eslesir(ORNEKLER[olay.mesaj]), olay.mesaj
    # Ayni onekli ama tutarsiz mesaj gizlenmez (isaret ayirt eder).
    assert detay_maskele(TUTARSIZ, TUTAR_OLAYLARI) == TUTARSIZ


def test_silme_motoru_sarmali_tutarli_metni_de_gizlenir() -> None:
    sarili = messages.deleted_with_dependents(
        ORNEKLER["progress_payment_deleted"], 2, ["Kalem 2"], []
    )
    assert detay_maskele(sarili, TUTAR_OLAYLARI) == GIZLI_TUTAR_METNI


def test_gizli_olaylar_kategoriden_secer() -> None:
    assert {o.mesaj for o in gizli_olaylar(frozenset({Hassas.maas_kisisel}))} == {
        "payroll_period_paid"
    }
    assert gizli_olaylar(frozenset({Hassas.banka_kasa})) == ()


# --- API: liste + Excel -------------------------------------------------------------------------


async def _satirlari_yaz(seeded_db) -> None:
    await seeded_db.execute(delete(AuditLog))
    simdi = datetime.now(UTC)
    for sira, detay in enumerate([*ORNEKLER.values(), TUTARSIZ]):
        seeded_db.add(
            AuditLog(
                action=AuditAction.update, detail=detay, occurred_at=simdi.replace(second=sira)
            )
        )
    await seeded_db.flush()


async def _giris(client, seeded_db, user_factory, anahtar: str, gizli: set[H]) -> dict[str, str]:
    rol = await rol_gizli(seeded_db, anahtar, gizli)
    e_posta = f"{anahtar}@b5a.co"
    await user_factory(email=e_posta, password=_PAROLA, role_key=rol.key)
    yanit = await client.post("/auth/login", json={"email": e_posta, "password": _PAROLA})
    assert yanit.status_code == 200
    return {"Authorization": f"Bearer {yanit.json()['access_token']}"}


@pytest.mark.parametrize("bayrak", list(BEKLENEN_GIZLI))
async def test_liste_gizli_kategoride_tutarli_metin_degisir_digerleri_kalir(
    client, seeded_db, user_factory, bayrak: H
) -> None:
    await _satirlari_yaz(seeded_db)
    basliklar = await _giris(client, seeded_db, user_factory, f"g_{bayrak.value}", {bayrak})
    await seeded_db.execute(delete(AuditLog).where(AuditLog.detail.like("Sisteme giriş%")))

    yanit = await client.get("/audit-log?limit=100", headers=basliklar)
    assert yanit.status_code == 200
    metinler = [o["detail"] for o in yanit.json()["items"] if o["action"] == "update"]

    gizli_beklenen = BEKLENEN_GIZLI[bayrak]
    for anahtar, ornek in ORNEKLER.items():
        if anahtar in gizli_beklenen:
            assert ornek not in metinler, f"{bayrak.value}: {anahtar} metni SIZIYOR"
        else:
            assert ornek in metinler, f"{bayrak.value}: {anahtar} gereksiz gizlendi"
    assert metinler.count(GIZLI_TUTAR_METNI) == len(gizli_beklenen)
    assert TUTARSIZ in metinler  # tutarsiz kardes kalir


async def test_liste_kategorisi_acik_rol_tam_metni_gorur_ve_tur_aktor_zaman_kalir(
    client, seeded_db, user_factory
) -> None:
    await _satirlari_yaz(seeded_db)
    basliklar = await _giris(client, seeded_db, user_factory, "acik_rol", set())
    await seeded_db.execute(delete(AuditLog).where(AuditLog.detail.like("Sisteme giriş%")))
    yanit = await client.get("/audit-log?limit=100", headers=basliklar)
    ogeler = [o for o in yanit.json()["items"] if o["action"] == "update"]
    assert {o["detail"] for o in ogeler} == {*ORNEKLER.values(), TUTARSIZ}


async def test_liste_gizli_rolde_tur_ve_zaman_korunur(client, seeded_db, user_factory) -> None:
    await _satirlari_yaz(seeded_db)
    basliklar = await _giris(client, seeded_db, user_factory, "hepsi", {H.tum_tutarlar})
    yanit = await client.get("/audit-log?limit=100&action=update", headers=basliklar)
    ogeler = yanit.json()["items"]
    gizli = [o for o in ogeler if o["detail"] == GIZLI_TUTAR_METNI]
    assert len(gizli) == len(ORNEKLER)
    assert all(o["action"] == "update" and o["occurred_at"] for o in gizli)


async def test_arama_gizli_tutari_kesfetmeye_yaramaz(client, seeded_db, user_factory) -> None:
    """`q` metin icinde arar: gizli okuyucu tutari 'tahmin ederek' dogrulayamamali."""
    await _satirlari_yaz(seeded_db)
    gizli = await _giris(client, seeded_db, user_factory, "arama_gizli", {H.maas_kisisel})
    acik = await _giris(client, seeded_db, user_factory, "arama_acik", set())
    assert (await client.get("/audit-log?q=1250.00", headers=acik)).json()["total"] >= 1
    yanit = (await client.get("/audit-log?q=Bordro dönemi ödendi", headers=gizli)).json()
    assert yanit["total"] == 0 and yanit["items"] == []
    # Hakedis (sozlesme_fiyat acik) bu rolde aranabilir kalir.
    assert (await client.get("/audit-log?q=Hakediş silindi", headers=gizli)).json()["total"] == 1


def _xlsx_detaylari(icerik: bytes) -> list[str]:
    sayfa = load_workbook(io.BytesIO(icerik)).active
    return [satir[3] for satir in sayfa.iter_rows(min_row=2, values_only=True)]


@pytest.mark.parametrize("bayrak", [H.maas_kisisel, H.sozlesme_fiyat, H.tum_tutarlar])
async def test_excel_ayni_kurala_uyar(client, seeded_db, user_factory, bayrak: H) -> None:
    await _satirlari_yaz(seeded_db)
    basliklar = await _giris(client, seeded_db, user_factory, f"x_{bayrak.value}", {bayrak})
    yanit = await client.get("/audit-log/export.xlsx", headers=basliklar)
    assert yanit.status_code == 200
    detaylar = _xlsx_detaylari(yanit.content)
    gizli_beklenen = BEKLENEN_GIZLI[bayrak]
    for anahtar, ornek in ORNEKLER.items():
        assert (ornek in detaylar) == (anahtar not in gizli_beklenen), (bayrak.value, anahtar)
    assert detaylar.count(GIZLI_TUTAR_METNI) == len(gizli_beklenen)


async def test_excel_kategorisi_acik_rol_tam_metni_alir(client, seeded_db, user_factory) -> None:
    await _satirlari_yaz(seeded_db)
    basliklar = await _giris(client, seeded_db, user_factory, "x_acik", set())
    yanit = await client.get("/audit-log/export.xlsx", headers=basliklar)
    detaylar = _xlsx_detaylari(yanit.content)
    assert set(ORNEKLER.values()) <= set(detaylar)
    assert GIZLI_TUTAR_METNI not in detaylar
