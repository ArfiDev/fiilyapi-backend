# ruff: noqa: F811  (admin/kab fiksturleri test_catalog_items_api.py'den ithal edilir)
"""KAT-B1.1 — onarim turu: kaynak kodu karakterleri (D2/D3), toplu yalniz-tarih (D5), hata
siralamasi (D7), Excel formul enjeksiyonu (D4), fiyat tarihi araligi (Madde 9), DB CHECK (Y1)."""

from __future__ import annotations

import uuid
from datetime import date
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from sqlalchemy.exc import IntegrityError

from app.modules.catalog import schemas
from app.modules.catalog.models import ContractorType, EvCatalogItem

from .test_catalog_items_api import (  # noqa: F401
    URL,
    _disiplin,
    _govde,
    admin,
    kab,
)

pytestmark = pytest.mark.asyncio

BULK = f"{URL}/bulk"
BUGUN = date(2026, 10, 3)


def _loc_listesi(resp) -> list[list]:
    return [e["loc"] for e in resp.json()["detail"]]


async def _toplu(client, admin, items, **ust):
    return await client.post(BULK, json={"items": items, **ust}, headers=admin)


# ------------------------------------------------------------ D2/D3 karakterler


@pytest.mark.parametrize("kod", ["a\x00b", "a​b", "​15.1", "a\tb", "a﻿b", "a‮b", "a\x07b"])
async def test_D3_kaynak_kodunda_kontrol_ve_bicim_karakteri_422_alan_adli(
    client, admin, kab, kod
) -> None:
    tekil = await client.post(URL, json=_govde(kab, source_code=kod), headers=admin)
    assert tekil.status_code == 422, tekil.text
    assert ["body", "source_code"] in _loc_listesi(tekil)
    toplu = await _toplu(
        client, admin, [_govde(kab, name="A"), _govde(kab, name="B", source_code=kod)]
    )
    assert toplu.status_code == 422, toplu.text
    assert ["body", "items", 1, "source_code"] in _loc_listesi(toplu)
    assert (await client.get(URL, headers=admin)).json()["items"] == []  # hicbiri yazilmadi


async def test_D3_PATCH_kaynak_kodu_kontrol_karakteri_422(client, admin, kab) -> None:
    item = (await client.post(URL, json=_govde(kab), headers=admin)).json()
    resp = await client.patch(f"{URL}/{item['id']}", json={"source_code": "a\x00b"}, headers=admin)
    assert resp.status_code == 422
    assert ["body", "source_code"] in _loc_listesi(resp)


async def test_D2_buyuk_kucuk_harf_duyarliligi_AYNEN_kalir(client, admin, kab) -> None:
    """Urun karari CEO'da: simdilik `Y.1` ve `y.1` AYRI kodlardir (davranis kilitlenir)."""
    a = await client.post(URL, json=_govde(kab, name="A", source_code="Y.1"), headers=admin)
    b = await client.post(URL, json=_govde(kab, name="B", source_code="y.1"), headers=admin)
    assert (a.status_code, b.status_code) == (201, 201)


# ------------------------------------------------------------ D5


async def _fiyatli(client, admin, kab, **over) -> dict:
    r = await client.post(
        URL,
        json=_govde(
            kab,
            source_code="S.1",
            ref_price="10.00",
            ref_price_date="2025-01-01",
            **over,
        ),
        headers=admin,
    )
    assert r.status_code == 201, r.text
    return r.json()


async def test_D5_update_price_fiyatsiz_yalniz_tarih_PATCH_ile_ayni_sonuc(
    client, admin, kab
) -> None:
    seed = await _fiyatli(client, admin, kab)
    resp = await _toplu(
        client,
        admin,
        [_govde(kab, name="x", source_code="S.1", ref_price_date="2026-01-01")],
        on_source_conflict="update_price",
    )
    assert resp.status_code == 201, resp.text
    assert resp.json()["items"][0]["action"] == "price_updated"
    now = (await client.get(URL, headers=admin)).json()["items"][0]
    assert (now["ref_price"], now["ref_price_date"]) == ("10.00", "2026-01-01")
    assert now["price_updated_at"] == seed["price_updated_at"]  # damga DEGISMEDI


async def test_D5_update_price_fiyatsiz_kalemde_yalniz_tarih_422(client, admin, kab) -> None:
    await client.post(URL, json=_govde(kab, source_code="S.1"), headers=admin)
    resp = await _toplu(
        client,
        admin,
        [_govde(kab, name="x", source_code="S.1", ref_price_date="2026-01-01")],
        on_source_conflict="update_price",
    )
    assert resp.status_code == 422, resp.text
    assert resp.json()["errors"][0]["loc"] == ["body", "items", 0, "ref_price_date"]
    now = (await client.get(URL, headers=admin)).json()["items"][0]
    assert now["ref_price_date"] is None


# ------------------------------------------------------------ D7


async def test_D7_hatalar_istek_indeksine_gore_SIRALI_girdi_siradan_bagimsiz(
    client, admin, kab
) -> None:
    await client.post(URL, json=_govde(kab, name="Beton"), headers=admin)
    resp = await _toplu(
        client,
        admin,
        [
            _govde(kab, name="beton"),  # 0: DB ad+birim cakismasi (SON asamada toplanir)
            {**_govde(kab, name="Yok"), "discipline_id": str(uuid.uuid4())},  # 1: ILK asamada
        ],
    )
    assert resp.status_code == 422, resp.text
    assert [e["loc"] for e in resp.json()["errors"]] == [
        ["body", "items", 0, "name"],
        ["body", "items", 1, "discipline_id"],
    ]


# ------------------------------------------------------------ D4


async def test_D4_excel_formul_enjeksiyonu_hucreler_string_yazilir(client, admin, kab) -> None:
    payload = '=HYPERLINK("http://e.vil","x")'
    ids = []
    for ad, kod, birim in [
        (payload, "=1+1", "@x"),
        ("+SUM(A1)", "-2+3", "\tsekme"),
        ("\r=cr", "@ref", "m"),
    ]:
        r = await client.post(
            URL, json=_govde(kab, name=ad, uom=birim, source_code=kod), headers=admin
        )
        assert r.status_code == 201, r.text
        ids.append(r.json())
    resp = await client.get(f"{URL}/export", headers=admin)
    assert resp.status_code == 200
    sayfa = openpyxl.load_workbook(BytesIO(resp.content)).active
    metin_hucreleri = [c for satir in sayfa.iter_rows(min_row=2) for c in satir if c.value]
    tehlikeli = [
        c for c in metin_hucreleri if str(c.value).startswith(("=", "+", "-", "@", "\t", "\r"))
    ]
    # ad/kod/birim hucreleri (ad/birim sema `strip`i bastaki sekme/CR'yi atar → 7)
    assert len(tehlikeli) == 7, [c.value for c in tehlikeli]
    assert {c.data_type for c in tehlikeli} == {"s"}, [(c.value, c.data_type) for c in tehlikeli]
    assert any(c.value == payload for c in tehlikeli)  # icerik AYNEN korundu


# ------------------------------------------------------------ Madde 9: fiyat tarihi araligi


@pytest.fixture
def bugun_sabit(monkeypatch):
    monkeypatch.setattr(schemas, "today", lambda: BUGUN)
    return BUGUN


def _yuksek(gun: date) -> date:
    return gun.fromordinal(gun.toordinal() + schemas.REF_PRICE_DATE_MAX_DAYS_AHEAD)


async def test_madde9_sinir_sabitleri_adli_ve_CEO_karariyla_ayni() -> None:
    assert schemas.REF_PRICE_DATE_MIN == date(2000, 1, 1)
    assert schemas.REF_PRICE_DATE_MAX_DAYS_AHEAD == 366


async def test_madde9_tekil_POST_iki_uc_dahil_bir_gun_disi_422(
    client, admin, kab, bugun_sabit
) -> None:
    alt, ust = schemas.REF_PRICE_DATE_MIN, _yuksek(bugun_sabit)
    one = date.fromordinal
    for i, (gun, beklenen) in enumerate(
        [
            (alt, 201),
            (ust, 201),
            (one(alt.toordinal() - 1), 422),
            (one(ust.toordinal() + 1), 422),
        ]
    ):
        r = await client.post(
            URL,
            json=_govde(kab, name=f"K{i}", ref_price="1.00", ref_price_date=gun.isoformat()),
            headers=admin,
        )
        assert r.status_code == beklenen, (gun, r.text)
        if beklenen == 422:
            assert ["body", "ref_price_date"] in _loc_listesi(r)


async def test_madde9_PATCH_ve_toplu_ayni_kurali_paylasir(client, admin, kab, bugun_sabit) -> None:
    item = (
        await client.post(URL, json=_govde(kab, ref_price="1.00", source_code="S.1"), headers=admin)
    ).json()
    ust = _yuksek(bugun_sabit)
    asan = date.fromordinal(ust.toordinal() + 1)
    asagi = date.fromordinal(schemas.REF_PRICE_DATE_MIN.toordinal() - 1)
    for gun in (asan, asagi):
        r = await client.patch(
            f"{URL}/{item['id']}", json={"ref_price_date": gun.isoformat()}, headers=admin
        )
        assert r.status_code == 422, r.text
        assert ["body", "ref_price_date"] in _loc_listesi(r)
        r = await _toplu(
            client,
            admin,
            [_govde(kab, name="Z", ref_price="2.00", ref_price_date=gun.isoformat())],
        )
        assert r.status_code == 422, r.text
        assert ["body", "items", 0, "ref_price_date"] in _loc_listesi(r)
    ok = await client.patch(
        f"{URL}/{item['id']}", json={"ref_price_date": ust.isoformat()}, headers=admin
    )
    assert ok.status_code == 200, ok.text
    ok2 = await _toplu(
        client,
        admin,
        [
            _govde(
                kab,
                name="Z",
                ref_price="2.00",
                ref_price_date=schemas.REF_PRICE_DATE_MIN.isoformat(),
            )
        ],
    )
    assert ok2.status_code == 201, ok2.text


# ------------------------------------------------------------ DB CHECK (Y1 yedegi)


async def test_Y1_DB_CHECK_fiyatsiz_fiyat_tarihini_reddeder(db_session, seeded_db) -> None:
    disiplin = await _disiplin(db_session, "CK1")
    satir = EvCatalogItem(
        discipline_id=disiplin.id,
        name="k",
        uom="m",
        standard_unit_mhr=Decimal("1"),
        default_contractor_type=ContractorType.OWN,
        poz_no="CK1-0001",
        ref_price=None,
        ref_price_date=date(2026, 1, 1),
    )
    with pytest.raises(IntegrityError) as exc:
        async with db_session.begin_nested():
            db_session.add(satir)
            await db_session.flush()
    assert "ck_ev_catalog_items_ref_price_date_requires_price" in str(exc.value.orig)


async def test_D4_yardimci_sekme_ve_CR_ile_baslayan_metni_de_string_yazar() -> None:
    from app.modules.catalog.export import _write

    sayfa = openpyxl.Workbook().active
    for satir, deger in enumerate(["\t=x", "\r=x", "=x", "+1", "-1", "@x", "normal"], start=1):
        _write(sayfa, satir, 1, deger)
    assert [sayfa.cell(row=r, column=1).data_type for r in range(1, 8)] == ["s"] * 7
