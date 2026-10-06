"""IZN-B5a — personel sızıntı onarımları (madde 23a, 21c).

23a: `assigned_project_id` / `assigned_section_id` aktörün GÖREMEDİĞİ projeye işaret ederse yanıt,
projenin hiç var olmamasıyla BİREBİR aynıdır (404, aynı gövde) — varlık sızmaz.
21c: tekillik denetlenen (`tc_no`) alan, aktörün maskesinde GİZLİyse POST'ta DOLU gönderilemez
(tekillik denetiminden ÖNCE 403): aksi hâlde 409 "TCKN kayıtlı" bilgisi sızardı.
"""

import uuid
from datetime import date

import pytest
from sqlalchemy import select

from app.modules.sites.models import Section, Site
from app.modules.users.models import User
from tests._proje_ekibi import ekibe_ekle

pytestmark = pytest.mark.asyncio

KAYITLI_TCKN = "10000000146"
AD = {"full_name": "Ali Veli", "source": "company"}


async def _login(client, user_factory, role_key: str, email: str) -> dict[str, str]:
    await user_factory(email=email, password="parola1234", role_key=role_key)
    resp = await client.post("/auth/login", json={"email": email, "password": "parola1234"})
    assert resp.status_code == 200, resp.text
    return {"Authorization": f"Bearer {resp.json()['access_token']}"}


@pytest.fixture
async def dunya(client, seeded_db, user_factory, project_factory):
    """İK kullanıcısı YALNIZ A projesinin ekibinde; B projesi ona görünmez."""
    proje_a = await project_factory("B5A-A", name="Görünen")
    proje_b = await project_factory("B5A-B", name="Görünmeyen")
    sahibi = {}
    for proje in (proje_a, proje_b):
        site = Site(project_id=proje.id, code=f"{proje.code}-S", name=f"{proje.code} Şantiye")
        seeded_db.add(site)
        await seeded_db.flush()
        bolum = Section(
            site_id=site.id,
            name=f"{proje.code} Bölüm",
            start_date=date(2026, 5, 4),
            end_date=date(2026, 5, 29),
            planned_worker_count=3,
            sort_order=1,
        )
        seeded_db.add(bolum)
        await seeded_db.flush()
        sahibi[proje.code] = bolum
    basliklar = await _login(client, user_factory, "hr_manager", "ik-b5a@personnel.co")
    kullanici = (
        await seeded_db.execute(select(User).where(User.email == "ik-b5a@personnel.co"))
    ).scalar_one()
    await ekibe_ekle(seeded_db, kullanici, proje_a.id)
    admin = await _login(client, user_factory, "system_admin", "admin-b5a@personnel.co")
    return {"a": proje_a, "b": proje_b, "bolum": sahibi, "ik": basliklar, "admin": admin}


async def _kayitsiz_404_govdesi(client, basliklar) -> dict:
    yanit = await client.post(
        "/personnel",
        json={**AD, "assigned_project_id": str(uuid.uuid4())},
        headers=basliklar,
    )
    assert yanit.status_code == 404, yanit.text
    return yanit.json()


async def test_gorunmeyen_projeye_atama_POST_404_olmayan_projeyle_ayni_govde(client, dunya):
    beklenen = await _kayitsiz_404_govdesi(client, dunya["ik"])
    yanit = await client.post(
        "/personnel",
        json={**AD, "assigned_project_id": str(dunya["b"].id)},
        headers=dunya["ik"],
    )
    assert yanit.status_code == 404, yanit.text
    assert yanit.json() == beklenen


async def test_gorunen_projeye_atama_POST_201(client, dunya):
    yanit = await client.post(
        "/personnel",
        json={**AD, "assigned_project_id": str(dunya["a"].id)},
        headers=dunya["ik"],
    )
    assert yanit.status_code == 201, yanit.text


async def test_sistem_yoneticisi_her_projeye_atar(client, dunya):
    yanit = await client.post(
        "/personnel",
        json={**AD, "assigned_project_id": str(dunya["b"].id)},
        headers=dunya["admin"],
    )
    assert yanit.status_code == 201, yanit.text


async def test_gorunmeyen_projeye_atama_PATCH_404_ve_kayit_degismez(client, dunya):
    olustur = await client.post("/personnel", json=AD, headers=dunya["ik"])
    assert olustur.status_code == 201, olustur.text
    kimlik = olustur.json()["id"]
    yanit = await client.patch(
        f"/personnel/{kimlik}",
        json={"assigned_project_id": str(dunya["b"].id)},
        headers=dunya["ik"],
    )
    assert yanit.status_code == 404, yanit.text
    sonra = (await client.get(f"/personnel/{kimlik}", headers=dunya["ik"])).json()
    assert sonra["assigned_project_id"] is None


async def test_gorunmeyen_projenin_bolumu_404_gorunen_projeyle_birlikte_bile(client, dunya):
    """Bölüm B'nin, proje A (görünen) gönderilse de: varlık sızmasın → 404 (422 değil)."""
    olmayan = await client.post(
        "/personnel",
        json={
            **AD,
            "assigned_project_id": str(dunya["a"].id),
            "assigned_section_id": str(uuid.uuid4()),
        },
        headers=dunya["ik"],
    )
    assert olmayan.status_code == 404, olmayan.text
    yanit = await client.post(
        "/personnel",
        json={
            **AD,
            "assigned_project_id": str(dunya["a"].id),
            "assigned_section_id": str(dunya["bolum"]["B5A-B"].id),
        },
        headers=dunya["ik"],
    )
    assert yanit.status_code == 404, yanit.text
    assert yanit.json() == olmayan.json()


async def test_gorunen_projenin_bolumu_POST_201(client, dunya):
    yanit = await client.post(
        "/personnel",
        json={
            **AD,
            "assigned_project_id": str(dunya["a"].id),
            "assigned_section_id": str(dunya["bolum"]["B5A-A"].id),
        },
        headers=dunya["ik"],
    )
    assert yanit.status_code == 201, yanit.text


# --- 21c ----------------------------------------------------------------------------------------


@pytest.fixture
async def kayitli_tckn(client, admin_headers):
    yanit = await client.post(
        "/personnel", json={**AD, "tc_no": KAYITLI_TCKN}, headers=admin_headers
    )
    assert yanit.status_code == 201, yanit.text


async def test_gizli_rol_kayitli_tckn_ile_POST_403_409_degil(client, gizli_headers, kayitli_tckn):
    yanit = await client.post(
        "/personnel", json={**AD, "full_name": "Veli", "tc_no": KAYITLI_TCKN}, headers=gizli_headers
    )
    assert yanit.status_code == 403, yanit.text
    assert "tc_no" in yanit.json()["detail"]


async def test_gizli_rol_kayitsiz_TCKN_ile_de_POST_403(client, gizli_headers):
    """Doluluk tek başına yetmez/gerekir: tekillik sonucu ayrışmasın diye HER dolu gönderim 403."""
    yanit = await client.post(
        "/personnel", json={**AD, "tc_no": "10000000146"}, headers=gizli_headers
    )
    assert yanit.status_code == 403, yanit.text


async def test_gizli_rol_PATCH_ile_ayni_govde(client, gizli_headers, admin_headers, kayitli_tckn):
    kisi = await client.post("/personnel", json={**AD, "full_name": "Z"}, headers=admin_headers)
    patch = await client.patch(
        f"/personnel/{kisi.json()['id']}", json={"tc_no": KAYITLI_TCKN}, headers=gizli_headers
    )
    post = await client.post(
        "/personnel", json={**AD, "tc_no": KAYITLI_TCKN}, headers=gizli_headers
    )
    assert patch.status_code == post.status_code == 403
    assert patch.json() == post.json()


async def test_gizli_rol_tc_no_bos_birakarak_personel_olusturur(client, gizli_headers):
    yanit = await client.post("/personnel", json=AD, headers=gizli_headers)
    assert yanit.status_code == 201, yanit.text
    assert yanit.json()["tc_no"] is None


async def test_gizli_rol_tekil_olmayan_gizli_alani_POST_edebilir(client, gizli_headers):
    """Daralma YALNIZ tekillik denetlenen alanlar içindir: iban/phone POST'ta serbest kalır."""
    yanit = await client.post(
        "/personnel",
        json={**AD, "phone": "05320000000", "iban": "TR330006100519786457841326"},
        headers=gizli_headers,
    )
    assert yanit.status_code == 201, yanit.text


async def test_gizlemeyen_rol_kayitli_tckn_409_alir(client, ik_headers, kayitli_tckn):
    yanit = await client.post("/personnel", json={**AD, "tc_no": KAYITLI_TCKN}, headers=ik_headers)
    assert yanit.status_code == 409, yanit.text


# --- 23a okuma tarafı ---------------------------------------------------------------------------


def _yayin(**fark) -> dict:
    return {
        "full_name": "Okuma Sızıntı",
        "source": "company",
        "tc_no": "10000000146",
        "birth_date": "1990-01-01",
        "phone": "5551112233",
        "address": "Mahalle Sokak No 1",
        "emergency_contact_name": "Ayşe",
        "emergency_contact_phone": "5559998877",
        "trade": "Kalıpçı",
        "hire_date": "2026-01-01",
        "wage_type": "daily",
        "wage_amount": "1500.00",
        "is_draft": False,
        **fark,
    }


@pytest.fixture
async def atanmislar(client, dunya):
    """Admin: biri görünmeyen B, biri görünen A projesine atanmış YAYINDA personel + eski belge."""
    sonuc = {}
    for ad, proje, tc in (
        ("Bde", dunya["b"], "10000000146"),
        ("Ade", dunya["a"], "11111111110"),
    ):
        yanit = await client.post(
            "/personnel",
            json=_yayin(full_name=ad, tc_no=tc, assigned_project_id=str(proje.id)),
            headers=dunya["admin"],
        )
        assert yanit.status_code == 201, yanit.text
        kimlik = yanit.json()["id"]
        belge = await client.post(
            f"/personnel/{kimlik}/documents",
            json={"free_label": "Ehliyet", "valid_until": "2020-01-01"},
            headers=dunya["admin"],
        )
        assert belge.status_code == 201, belge.text
        sonuc[ad] = kimlik
    return sonuc


async def test_liste_gorunmeyen_projeye_atama_null_gorunen_dolu(client, dunya, atanmislar):
    yanit = await client.get("/personnel", headers=dunya["ik"])
    satir = {s["full_name"]: s for s in yanit.json()["items"]}
    assert satir["Bde"]["assigned_project_id"] is None
    assert satir["Ade"]["assigned_project_id"] == str(dunya["a"].id)


async def test_admin_atamayi_gorur(client, dunya, atanmislar):
    yanit = await client.get("/personnel", headers=dunya["admin"])
    satir = {s["full_name"]: s for s in yanit.json()["items"]}
    assert satir["Bde"]["assigned_project_id"] == str(dunya["b"].id)


async def test_detay_ve_patch_yanitinda_gorunmeyen_atama_null(client, dunya, atanmislar):
    detay = await client.get(f"/personnel/{atanmislar['Bde']}", headers=dunya["ik"])
    assert detay.json()["assigned_project_id"] is None
    assert detay.json()["assigned_section_id"] is None
    patch = await client.patch(
        f"/personnel/{atanmislar['Bde']}", json={"trade": "Demirci"}, headers=dunya["ik"]
    )
    assert patch.status_code == 200, patch.text
    assert patch.json()["assigned_project_id"] is None
    ok = await client.get(f"/personnel/{atanmislar['Ade']}", headers=dunya["ik"])
    assert ok.json()["assigned_project_id"] == str(dunya["a"].id)


async def test_project_id_suzgeci_gorunmeyen_projede_bos_200_ayni_olmayanla(
    client, dunya, atanmislar
):
    beklenen = (
        await client.get(
            "/personnel", params={"project_id": str(uuid.uuid4())}, headers=dunya["ik"]
        )
    ).json()
    yanit = await client.get(
        "/personnel", params={"project_id": str(dunya["b"].id)}, headers=dunya["ik"]
    )
    assert yanit.status_code == 200
    assert yanit.json() == beklenen
    assert yanit.json()["items"] == [] and yanit.json()["total"] == 0
    gorunen = await client.get(
        "/personnel", params={"project_id": str(dunya["a"].id)}, headers=dunya["ik"]
    )
    assert gorunen.json()["total"] == 1
    admin = await client.get(
        "/personnel", params={"project_id": str(dunya["b"].id)}, headers=dunya["admin"]
    )
    assert admin.json()["total"] == 1


async def test_export_gorunmeyen_proje_suzgeci_bos(client, dunya, atanmislar):
    yanit = await client.get(
        "/personnel/export.xlsx", params={"project_id": str(dunya["b"].id)}, headers=dunya["ik"]
    )
    assert yanit.status_code == 200
    from io import BytesIO

    from openpyxl import load_workbook

    sayfa = load_workbook(BytesIO(yanit.content)).active
    assert "Bde" not in {c.value for satir in sayfa.iter_rows() for c in satir}


async def test_ozet_gorunmeyen_projenin_adini_basmaz(client, dunya, atanmislar):
    ik = (await client.get("/hr/documents/summary", headers=dunya["ik"])).json()
    adlar = {s["personnel_name"]: s["project_name"] for s in ik["expired_documents"]}
    assert adlar["Bde"] is None
    assert adlar["Ade"] == dunya["a"].name
    admin = (await client.get("/hr/documents/summary", headers=dunya["admin"])).json()
    adlar = {s["personnel_name"]: s["project_name"] for s in admin["expired_documents"]}
    assert adlar["Bde"] == dunya["b"].name


async def test_izin_talebi_listesi_gorunmeyen_proje_suzgeci_bos(client, dunya, atanmislar):
    yanit = await client.get(
        "/leave-requests", params={"project_id": str(dunya["b"].id)}, headers=dunya["ik"]
    )
    assert yanit.status_code == 200
    assert yanit.json()["items"] == [] and yanit.json()["total"] == 0


# --- Görünmeyen projedeki atamayı PATCH ile silme/değiştirme = gizli projede yazma --------------


async def _atama(session, personel_id: str):
    from app.modules.personnel.models import Personnel

    kayit = await session.get(Personnel, uuid.UUID(personel_id))
    await session.refresh(kayit)
    return kayit.assigned_project_id


@pytest.mark.parametrize(
    "govde",
    [
        {"assigned_project_id": None},
        {"assigned_section_id": None},
        {"assigned_project_id": "A", "assigned_section_id": None},
    ],
)
async def test_gorunmeyen_projedeki_atamayi_degistiren_patch_404_atama_korunur(
    client, dunya, atanmislar, seeded_db, govde
):
    if govde.get("assigned_project_id") == "A":
        govde = {**govde, "assigned_project_id": str(dunya["a"].id)}
    yanit = await client.patch(f"/personnel/{atanmislar['Bde']}", json=govde, headers=dunya["ik"])
    assert yanit.status_code == 404, yanit.text
    assert yanit.json()["detail"] == "Atanan proje bulunamadı"
    assert await _atama(seeded_db, atanmislar["Bde"]) == dunya["b"].id


async def test_atama_alani_olmayan_patch_gorunmeyen_atamayi_korur(
    client, dunya, atanmislar, seeded_db
):
    yanit = await client.patch(
        f"/personnel/{atanmislar['Bde']}", json={"phone": "5550001122"}, headers=dunya["ik"]
    )
    assert yanit.status_code == 200, yanit.text
    assert await _atama(seeded_db, atanmislar["Bde"]) == dunya["b"].id


async def test_admin_gorunmeyen_atamayi_degistirebilir(client, dunya, atanmislar, seeded_db):
    yanit = await client.patch(
        f"/personnel/{atanmislar['Bde']}",
        json={"assigned_project_id": str(dunya["a"].id)},
        headers=dunya["admin"],
    )
    assert yanit.status_code == 200, yanit.text
    assert await _atama(seeded_db, atanmislar["Bde"]) == dunya["a"].id
