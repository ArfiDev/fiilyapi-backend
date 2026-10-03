"""Fiyatli Is Kalemi Katalogu Excel ciktisi (TKL-B5.2).

Saf sunum katmani: `WorkItemRead` listesini (kapsam maskesinden gecmis) calisma kitabina cevirir.
FLOAT-YASAK (boq/export.py emsali): her hucre `str`; maskeli/bos deger (`None`) hucreye
YAZILMAZ. Tarihler goruntuleme saat diliminde `gg.aa.yyyy`.
"""

from __future__ import annotations

from collections.abc import Sequence
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from app.core.timezone import to_display
from app.modules.catalog.schemas import WorkItemRead

SHEET_TITLE = "İş Kalemi Kataloğu"
#: Indirilen dosya adi (mockup: "İş Kalemi Kataloğu").
FILENAME = "İş Kalemi Kataloğu.xlsx"
DATE_FORMAT = "%d.%m.%Y"

#: Sutun basliklari — sira ve metin degistirilmez.
COLUMN_HEADERS: tuple[str, ...] = (
    "Poz No",
    "Disiplin",
    "İş Kalemi",
    "Birim",
    "Referans Fiyat",
    "Fiyat Güncelleme",
    "Son Fiyat",
    "Kaynak",
    "Belge",
    "Tarih",
    "Adam-saat/birim",
    "Varsayılan Yüklenici",
    # KAT-B1: yeni sutunlar SONA eklenir (eskilerin sirasi/metni degismez).
    "Kaynak Poz No",
    "Fiyat Tarihi",
)
_CONTRACTOR_LABELS = {"own": "Kendi", "subcon": "Taşeron"}
_COLUMN_WIDTHS = (12, 18, 40, 8, 16, 16, 14, 10, 24, 12, 16, 20, 16, 14)


def _day(value: datetime | None) -> str | None:
    return None if value is None else to_display(value).strftime(DATE_FORMAT)


def _date(value: date | None) -> str | None:
    """`Fiyat Tarihi` takvim gunudur (saat dilimi donusumu YOK): `gg.aa.yyyy`."""
    return None if value is None else value.strftime(DATE_FORMAT)


#: Excel'de hucreyi FORMUL yapan ilk karakterler (CSV/formul enjeksiyonu); sekme ve CR dahil.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _write(sheet, row: int, column: int, value: str) -> None:  # noqa: ANN001
    """Metin hucresini HER ZAMAN string yazar (`data_type='s'`): `=HYPERLINK(...)` gibi bir
    ad/kod formul olarak calismaz (KAT-B1.1 D4)."""
    cell = sheet.cell(row=row, column=column)
    cell.value = value
    if value.startswith(_FORMULA_PREFIXES):
        cell.data_type = "s"


def _s(value: object | None) -> str | None:
    return None if value is None else str(value)


def _row(item: WorkItemRead) -> tuple[str | None, ...]:
    last = item.last_price
    return (
        item.poz_no,
        item.discipline.name,
        item.name,
        item.uom,
        _s(item.ref_price),
        _day(item.price_updated_at),
        _s(last.price) if last else None,
        last.source if last else None,
        last.doc_no if last else None,
        _day(last.at) if last else None,
        _s(item.standard_unit_mhr),
        _CONTRACTOR_LABELS.get(item.default_contractor_type, item.default_contractor_type),
        item.source_code,
        _date(item.ref_price_date),
    )


def build_catalog_workbook(items: Sequence[WorkItemRead]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_TITLE
    for column, header in enumerate(COLUMN_HEADERS, start=1):
        cell = sheet.cell(row=1, column=column)
        cell.value = header
        cell.font = Font(bold=True)
        sheet.column_dimensions[cell.column_letter].width = _COLUMN_WIDTHS[column - 1]
    for row, item in enumerate(items, start=2):
        for column, value in enumerate(_row(item), start=1):
            if value is not None:
                _write(sheet, row, column, value)
    sheet.freeze_panes = "A2"
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
